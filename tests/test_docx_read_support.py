"""Focused coverage for semantic, bounded DOCX reads."""

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from squid_core.file_readers.docx import read_docx
from squid_core.file_readers.registry import reader_for_suffix
from squid_core.settings import ReadFileSettings
from squid_core.tools.builtins.read_file import execute as read_file


def test_registry_dispatches_docx_and_preserves_semantic_blocks(
    tmp_path: Path,
) -> None:
    source = tmp_path / "brief.docx"
    document = Document()
    document.add_heading("Project brief", level=1)
    document.add_paragraph("Opening paragraph.")
    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("First numbered item", style="List Number")
    _add_hyperlink(
        document.add_paragraph("Visit "),
        "Code Squid",
        "https://example.test",
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "one | two"
    document.add_paragraph("Closing paragraph.")
    document.save(source)

    result = read_file(tmp_path, str(source))

    assert reader_for_suffix(".docx") is read_docx
    assert reader_for_suffix(".DOCX") is read_docx
    assert result["ok"] is True
    assert result["media_type"] == "document"
    assert result["format"] == "docx"
    assert result["content"] == (
        "# Project brief\n"
        "Opening paragraph.\n"
        "- First bullet\n"
        "1. First numbered item\n"
        "Visit [Code Squid](https://example.test)\n"
        "| Name | Value |\n"
        "| --- | --- |\n"
        "| alpha | one \\| two |\n"
        "Closing paragraph.\n"
    )


def test_docx_preserves_paragraph_table_paragraph_order(tmp_path: Path) -> None:
    source = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "middle"
    document.add_paragraph("after")
    document.save(source)

    content = read_file(tmp_path, str(source))["content"]

    assert content.index("before") < content.index("middle") < content.index("after")


def test_docx_uses_text_truncation_and_targeted_line_ranges(tmp_path: Path) -> None:
    source = tmp_path / "long.docx"
    document = Document()
    for number in range(1, 21):
        document.add_paragraph(f"paragraph-{number}")
    document.save(source)
    settings = ReadFileSettings(default_max_lines=6, minimum_range_lines=1)

    whole = read_file(tmp_path, str(source), settings=settings)
    targeted = read_file(
        tmp_path,
        str(source),
        ranges=[{"start_line": 10, "end_line": 11}],
        settings=settings,
    )

    assert whole["content_truncated"] is True
    assert whole["total_lines"] == 20
    assert whole["returned_line_ranges"] == [
        {"start_line": 1, "end_line": 3},
        {"start_line": 18, "end_line": 20},
    ]
    assert "Use ranges" in whole["truncation_guidance"]
    assert "paragraph-10" not in whole["content"]
    assert targeted["content"] == "paragraph-10\nparagraph-11\n"
    assert targeted["returned_line_ranges"] == [
        {"start_line": 10, "end_line": 11}
    ]
    assert targeted["ranged"] is True


def test_docx_rejects_malformed_and_non_docx_containers(tmp_path: Path) -> None:
    malformed = tmp_path / "malformed.docx"
    malformed.write_bytes(b"not a zip")
    not_docx = tmp_path / "not-docx.docx"
    with ZipFile(not_docx, "w", ZIP_DEFLATED) as archive:
        archive.writestr("content.txt", "ordinary zip")

    malformed_result = read_file(tmp_path, str(malformed))
    parser_result = read_file(tmp_path, str(not_docx))

    assert malformed_result == {
        "ok": False,
        "tool": "read_file",
        "path": str(malformed),
        "error": "invalid or unreadable DOCX container",
    }
    assert parser_result == {
        "ok": False,
        "tool": "read_file",
        "path": str(not_docx),
        "error": "invalid or unreadable DOCX document",
    }


def test_docx_rejects_oversized_compressed_container(
    monkeypatch,
    tmp_path: Path,
) -> None:
    from squid_core.file_readers import docx

    source = tmp_path / "large.docx"
    Document().save(source)
    monkeypatch.setattr(docx, "MAX_DOCX_BYTES", source.stat().st_size - 1)

    result = read_file(tmp_path, str(source))

    assert result["ok"] is False
    assert "maximum supported size" in result["error"]
    assert "expanded DOCX" not in result["error"]


def test_docx_rejects_excessive_expanded_container(
    monkeypatch,
    tmp_path: Path,
) -> None:
    from squid_core.file_readers import docx

    source = tmp_path / "expanded.docx"
    Document().save(source)
    monkeypatch.setattr(docx, "MAX_DOCX_EXPANDED_BYTES", 2)

    result = read_file(tmp_path, str(source))

    assert result["ok"] is False
    assert "expanded DOCX" in result["error"]
    assert "maximum supported size is 2 bytes" in result["error"]


def test_docx_rejects_non_line_selectors(tmp_path: Path) -> None:
    source = tmp_path / "selectors.docx"
    Document().save(source)

    pages = read_file(
        tmp_path,
        str(source),
        pages=[{"start_page": 1, "end_page": 1}],
    )
    sheet = read_file(tmp_path, str(source), sheet="Sheet1")
    rows = read_file(
        tmp_path,
        str(source),
        sheet="Sheet1",
        rows={"start_row": 1, "end_row": 2},
    )

    assert pages["error"] == "page ranges are supported only for PDF files"
    assert sheet["error"] == "sheet and rows are supported only for XLSX and ODS files"
    assert rows["error"] == "sheet and rows are supported only for XLSX and ODS files"


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    """Add one external hyperlink using python-docx's supported relationship API."""

    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
