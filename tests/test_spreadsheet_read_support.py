"""Focused coverage for formula-preserving XLSX and ODS reads."""

from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table, TableCell, TableRow
from odf.text import P
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from squid_core.file_readers.ods import read_ods
from squid_core.file_readers.registry import reader_for_suffix
from squid_core.file_readers.xlsx import read_xlsx
from squid_core.settings import ReadFileSettings
from squid_core.tools.builtins.read_file import execute as read_file


def test_registry_dispatches_xlsx_and_preserves_sheet_values_and_formulas(
    tmp_path: Path,
) -> None:
    source = tmp_path / "budget.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Cost"])
    sheet.append(["Hosting", 12])
    sheet["B3"] = "=SUM(B2:B2)"
    workbook.create_sheet("Empty")
    workbook.save(source)

    result = read_file(tmp_path, str(source))

    assert reader_for_suffix(".xlsx") is read_xlsx
    assert result["ok"] is True
    assert result["media_type"] == "spreadsheet"
    assert result["format"] == "xlsx"
    assert result["sheet_names"] == ["Budget", "Empty"]
    assert "<<< SHEET: Budget >>>" in result["content"]
    assert "Item\tCost" in result["content"]
    assert "Hosting\t12" in result["content"]
    assert "[formula: =SUM(B2:B2)]" in result["content"]
    assert "<<< SHEET: Empty >>>" in result["content"]
    assert "[empty sheet]" in result["content"]


def test_registry_dispatches_ods_and_preserves_sheet_values_and_formulas(
    tmp_path: Path,
) -> None:
    source = tmp_path / "forecast.ods"
    document = OpenDocumentSpreadsheet()
    data = Table(name="Forecast")
    data.addElement(_ods_row("Month", "Revenue"))
    data.addElement(_ods_row("January", 20))
    formula_row = TableRow()
    formula_row.addElement(TableCell())
    formula = TableCell(
        valuetype="float",
        value=20,
        formula="of:=SUM([.B2:.B2])",
    )
    formula.addElement(P(text="20"))
    formula_row.addElement(formula)
    data.addElement(formula_row)
    document.spreadsheet.addElement(data)
    document.spreadsheet.addElement(Table(name="Empty"))
    document.save(str(source))

    result = read_file(tmp_path, str(source))

    assert reader_for_suffix(".ODS") is read_ods
    assert result["ok"] is True
    assert result["format"] == "ods"
    assert result["sheet_names"] == ["Forecast", "Empty"]
    assert "Month\tRevenue" in result["content"]
    assert "January\t20" in result["content"]
    assert "20 [formula: of:=SUM([.B2:.B2])]" in result["content"]
    assert "<<< SHEET: Empty >>>" in result["content"]
    assert "[empty sheet]" in result["content"]


def test_xlsx_chart_is_ignored_without_crashing_read(tmp_path: Path) -> None:
    source = tmp_path / "chart.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chart Data"
    sheet.append(["Quarter", "Sales"])
    sheet.append(["Q1", 8])
    chart = BarChart()
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=2),
        titles_from_data=True,
    )
    sheet.add_chart(chart, "D2")
    workbook.save(source)

    result = read_file(tmp_path, str(source))

    assert result["ok"] is True
    assert "Q1\t8" in result["content"]


def test_spreadsheet_reads_reject_text_ranges(tmp_path: Path) -> None:
    source = tmp_path / "simple.xlsx"
    workbook = Workbook()
    workbook.save(source)

    result = read_file(
        tmp_path,
        str(source),
        ranges=[{"start_line": 1, "end_line": 2}],
    )

    assert result["ok"] is False
    assert "line ranges are not supported for spreadsheet files" in result["error"]


def test_spreadsheet_text_uses_full_read_head_tail_truncation(tmp_path: Path) -> None:
    source = tmp_path / "long.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Long"
    for number in range(20):
        sheet.append([f"value-{number}"])
    workbook.save(source)

    result = read_file(
        tmp_path,
        str(source),
        settings=ReadFileSettings(default_max_lines=10, minimum_range_lines=1),
    )

    assert result["ok"] is True
    assert result["content_truncated"] is True
    assert result["returned_line_ranges"][0] == {"start_line": 1, "end_line": 5}
    assert result["returned_line_ranges"][1]["end_line"] == result["total_lines"]
    assert "omitted from this spreadsheet read" in result["content"]
    assert "value-19" in result["content"]


def test_xlsx_can_select_one_sheet(tmp_path: Path) -> None:
    source = tmp_path / "departments.xlsx"
    workbook = Workbook()
    workbook.active.title = "North"
    workbook.active.append(["north-only"])
    workbook.create_sheet("South").append(["south-only"])
    workbook.save(source)

    result = read_file(tmp_path, str(source), sheet="South")

    assert result["ok"] is True
    assert result["selected_sheet"] == "South"
    assert result["sheet_names"] == ["South"]
    assert "south-only" in result["content"]
    assert "north-only" not in result["content"]


def test_ods_can_select_one_sheet(tmp_path: Path) -> None:
    source = tmp_path / "departments.ods"
    document = OpenDocumentSpreadsheet()
    north = Table(name="North")
    north.addElement(_ods_row("north-only"))
    south = Table(name="South")
    south.addElement(_ods_row("south-only"))
    document.spreadsheet.addElement(north)
    document.spreadsheet.addElement(south)
    document.save(str(source))

    result = read_file(tmp_path, str(source), sheet="South")

    assert result["ok"] is True
    assert result["selected_sheet"] == "South"
    assert "south-only" in result["content"]
    assert "north-only" not in result["content"]


def test_xlsx_row_selection_is_source_based_and_inclusive(tmp_path: Path) -> None:
    source = tmp_path / "rows.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for number in range(1, 7):
        sheet.cell(number, 2, f"row-{number}")
    workbook.save(source)

    result = read_file(
        tmp_path,
        str(source),
        sheet="Data",
        rows={"start_row": 3, "end_row": 4},
    )

    assert result["ok"] is True
    assert result["selected_row_range"] == {"start_row": 3, "end_row": 4}
    assert "Selected source rows: 3-4 (inclusive)" in result["content"]
    assert "Used range: B1:B6" in result["content"]
    assert "Returned range: B3:B4" in result["content"]
    assert "row-3" in result["content"]
    assert "row-4" in result["content"]
    assert "row-2" not in result["content"]
    assert "row-5" not in result["content"]


def test_ods_row_selection_is_source_based_and_inclusive(tmp_path: Path) -> None:
    source = tmp_path / "rows.ods"
    document = OpenDocumentSpreadsheet()
    data = Table(name="Data")
    for number in range(1, 7):
        data.addElement(_ods_row(f"row-{number}"))
    document.spreadsheet.addElement(data)
    document.save(str(source))

    result = read_file(
        tmp_path,
        str(source),
        sheet="Data",
        rows={"start_row": 3, "end_row": 4},
    )

    assert result["ok"] is True
    assert "Returned range: A3:A4" in result["content"]
    assert "row-3" in result["content"]
    assert "row-4" in result["content"]
    assert "row-2" not in result["content"]
    assert "row-5" not in result["content"]


def test_unknown_sheet_returns_structured_error(tmp_path: Path) -> None:
    source = tmp_path / "simple.xlsx"
    workbook = Workbook()
    workbook.active.title = "Known"
    workbook.save(source)

    result = read_file(tmp_path, str(source), sheet="Missing")

    assert result == {
        "ok": False,
        "tool": "read_file",
        "path": str(source),
        "error": "unknown spreadsheet sheet: 'Missing'",
    }


def test_spreadsheet_selection_arguments_are_validated(tmp_path: Path) -> None:
    source = tmp_path / "simple.xlsx"
    workbook = Workbook()
    workbook.save(source)

    assert read_file(
        tmp_path,
        str(source),
        rows={"start_row": 1, "end_row": 2},
    )["error"] == "rows requires a selected sheet"
    assert read_file(
        tmp_path,
        str(source),
        sheet="Sheet",
        rows={"start_row": 0, "end_row": 2},
    )["error"] == "rows: start_row must be >= 1"
    assert read_file(
        tmp_path,
        str(source),
        sheet="Sheet",
        rows={"start_row": 3, "end_row": 2},
    )["error"] == "rows: end_row must be >= start_row"


def test_spreadsheet_arguments_are_rejected_for_text_files(tmp_path: Path) -> None:
    source = tmp_path / "notes.txt"
    source.write_text("hello\n", encoding="utf-8")

    result = read_file(tmp_path, str(source), sheet="Sheet1")

    assert result["ok"] is False
    assert result["error"] == "sheet and rows are supported only for XLSX and ODS files"


def test_targeted_rows_recover_content_omitted_from_large_workbook(tmp_path: Path) -> None:
    source = tmp_path / "long.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Long"
    for number in range(1, 41):
        sheet.append([f"value-{number}"])
    workbook.save(source)
    settings = ReadFileSettings(default_max_lines=12, minimum_range_lines=1)

    whole = read_file(tmp_path, str(source), settings=settings)
    targeted = read_file(
        tmp_path,
        str(source),
        sheet="Long",
        rows={"start_row": 20, "end_row": 21},
        settings=settings,
    )

    assert whole["content_truncated"] is True
    assert "value-20" not in whole["content"]
    assert targeted["content_truncated"] is False
    assert "value-20" in targeted["content"]
    assert "value-21" in targeted["content"]

    broad_targeted = read_file(
        tmp_path,
        str(source),
        sheet="Long",
        rows={"start_row": 1, "end_row": 40},
        settings=settings,
    )
    assert broad_targeted["content_truncated"] is True
    assert "narrower rows" in broad_targeted["truncation_guidance"]


def test_xlsx_formula_shows_cached_value_without_recalculation(tmp_path: Path) -> None:
    source = tmp_path / "cached-formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 50
    sheet["A2"] = "=A1*2"
    sheet["A3"] = "=A1*3"
    workbook.save(source)
    _set_xlsx_cached_number(source, "A2", "100")
    original_bytes = source.read_bytes()

    result = read_file(tmp_path, str(source))

    assert result["ok"] is True
    assert "100 [formula: =A1*2]" in result["content"]
    assert "[formula: =A1*3]" in result["content"]
    assert "None [formula:" not in result["content"]
    # The reader only consumes stored XML values; it never saves or recalculates.
    assert source.read_bytes() == original_bytes


def test_shared_container_safety_preserves_spreadsheet_failures(
    monkeypatch,
    tmp_path: Path,
) -> None:
    from squid_core.file_readers import spreadsheet

    malformed = tmp_path / "malformed.xlsx"
    malformed.write_bytes(b"not a zip")
    malformed_result = read_file(tmp_path, str(malformed))

    source = tmp_path / "bounded.xlsx"
    Workbook().save(source)
    monkeypatch.setattr(
        spreadsheet, "MAX_SPREADSHEET_BYTES", source.stat().st_size - 1
    )
    compressed_result = read_file(tmp_path, str(source))
    monkeypatch.setattr(spreadsheet, "MAX_SPREADSHEET_BYTES", 64 * 1024 * 1024)
    monkeypatch.setattr(spreadsheet, "MAX_SPREADSHEET_EXPANDED_BYTES", 2)
    expanded_result = read_file(tmp_path, str(source))

    assert "invalid or unreadable spreadsheet container" in malformed_result["error"]
    assert compressed_result["error"].startswith("spreadsheet is ")
    assert "maximum supported size" in compressed_result["error"]
    assert expanded_result["error"].startswith("expanded spreadsheet is ")
    assert "maximum supported size is 2 bytes" in expanded_result["error"]


def _ods_row(*values: object) -> TableRow:
    """Build a simple ODS row with portable displayed text and typed values."""

    row = TableRow()
    for value in values:
        if isinstance(value, (int, float)):
            cell = TableCell(valuetype="float", value=value)
        else:
            cell = TableCell(valuetype="string")
        cell.addElement(P(text=str(value)))
        row.addElement(cell)
    return row


def _set_xlsx_cached_number(source: Path, coordinate: str, value: str) -> None:
    """Inject the cached value openpyxl intentionally does not write itself."""

    replacement = source.with_suffix(".rewritten.xlsx")
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(source, "r") as input_archive, ZipFile(
        replacement, "w", ZIP_DEFLATED
    ) as output_archive:
        for member in input_archive.infolist():
            payload = input_archive.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(payload)
                cell = root.find(f".//{{{namespace}}}c[@r='{coordinate}']")
                assert cell is not None
                cached = cell.find(f"{{{namespace}}}v")
                if cached is None:
                    cached = ElementTree.SubElement(cell, f"{{{namespace}}}v")
                cached.text = value
                payload = ElementTree.tostring(
                    root, encoding="utf-8", xml_declaration=False
                )
            output_archive.writestr(member, payload)
    replacement.replace(source)
