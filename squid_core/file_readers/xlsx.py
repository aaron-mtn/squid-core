"""Read modern Excel workbooks into bounded, formula-preserving text."""

from .context import FileReadContext
from .spreadsheet import (
    MAX_SHEET_CELLS,
    SpreadsheetSheet,
    check_spreadsheet_size,
    reject_spreadsheet_ranges,
    select_spreadsheet_content,
    spreadsheet_error,
    spreadsheet_result,
)


def read_xlsx(context: FileReadContext) -> dict:
    """Read cell values and formulas from an XLSX workbook without recalculation."""

    validation_error = reject_spreadsheet_ranges(context) or check_spreadsheet_size(
        context
    )
    if validation_error is not None:
        return validation_error

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            context.file_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        try:
            cached_workbook = load_workbook(
                context.file_path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            try:
                cached_by_name = {
                    worksheet.title: worksheet
                    for worksheet in cached_workbook.worksheets
                }
                sheets = [
                    _read_sheet(worksheet, cached_by_name[worksheet.title])
                    for worksheet in workbook.worksheets
                ]
            finally:
                cached_workbook.close()
        finally:
            workbook.close()
    except Exception as error:
        # Parser exception types vary by corrupt ZIP/XML part. Keep all malformed
        # workbook failures inside the normal Tool-result boundary.
        return spreadsheet_error(
            context.path, f"invalid or unreadable XLSX spreadsheet: {error}"
        )
    selected = select_spreadsheet_content(context, sheets)
    if isinstance(selected, dict):
        return selected
    return spreadsheet_result(context, "xlsx", selected)


def _read_sheet(worksheet, cached_worksheet) -> SpreadsheetSheet:
    """Normalize the meaningful rectangular area of one openpyxl worksheet."""

    max_row = worksheet.max_row or 0
    max_column = worksheet.max_column or 0
    if max_row * max_column > MAX_SHEET_CELLS:
        raise ValueError(
            f"sheet {worksheet.title!r} declares {max_row * max_column} cells; "
            f"maximum supported sheet area is {MAX_SHEET_CELLS} cells"
        )

    populated: list[tuple[int, int, object]] = []
    cached_rows = cached_worksheet.iter_rows()
    for row in worksheet.iter_rows():
        cached_row = next(cached_rows)
        for cell, cached_cell in zip(row, cached_row):
            if cell.value is not None:
                value = (
                    _formula_value(cell.value, cached_cell.value)
                    if cell.data_type == "f"
                    else cell.value
                )
                populated.append((cell.row, cell.column, value))
    if not populated:
        return SpreadsheetSheet(worksheet.title, None, None, ())

    min_row = min(row for row, _, _ in populated)
    max_used_row = max(row for row, _, _ in populated)
    min_column = min(column for _, column, _ in populated)
    max_used_column = max(column for _, column, _ in populated)
    values = {(row, column): value for row, column, value in populated}
    rows = tuple(
        tuple(
            values.get((row, column))
            for column in range(min_column, max_used_column + 1)
        )
        for row in range(min_row, max_used_row + 1)
    )
    return SpreadsheetSheet(worksheet.title, min_row, min_column, rows)


def _formula_value(formula: object, cached_value: object) -> str:
    """Render a stored result when present without asking openpyxl to calculate."""

    if cached_value is None:
        return f"[formula: {formula}]"
    return f"{cached_value} [formula: {formula}]"
